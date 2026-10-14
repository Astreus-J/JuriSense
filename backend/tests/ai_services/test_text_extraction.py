"""docs/05-ai/document-intelligence.md: per-format text extraction + pagination mode."""

import io

import pytest
from docx import Document as DocxDocument

from app.ai_services.domain.errors import DocumentParsingError
from app.ai_services.infrastructure.text_extraction import extract_text
from app.core.config import get_settings

_MINIMAL_PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"


def test_pdf_extraction_is_native_pagination() -> None:
    # A corrupt/minimal PDF still parses as zero real pages worth of text via pypdf's
    # lenient reader in most cases; use a valid-structure single-page PDF instead so
    # extraction actually succeeds end to end.
    import pypdf

    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    result = extract_text(buf.getvalue(), "application/pdf")
    assert result.pagination_mode == "native"
    assert len(result.pages) == 1


def test_corrupt_pdf_raises_document_parsing_error() -> None:
    with pytest.raises(DocumentParsingError):
        extract_text(b"not a pdf at all", "application/pdf")


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_docx_extraction_is_synthetic_pagination_and_includes_table_rows() -> None:
    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("Clause 1: Liability shall be unlimited.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "SLA"
    table.rows[0].cells[1].text = "99.9%"
    doc.save(buf)

    result = extract_text(buf.getvalue(), DOCX_MIME)
    assert result.pagination_mode == "synthetic"
    full_text = "\n".join(result.pages)
    assert "Liability shall be unlimited" in full_text
    assert "[TABLE ROW]" in full_text
    assert "99.9%" in full_text


def test_docx_decompression_bomb_guard_rejects_an_oversized_declared_uncompressed_size(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Security audit (2026-10-12): a real zip bomb declares a huge uncompressed size
    while the compressed bytes stay tiny -- the guard reads that declared size from the
    zip's own metadata, so monkeypatching the limit down to something any ordinary DOCX
    already exceeds proves the check rejects BEFORE `DocxDocument` ever decompresses
    anything, without needing to actually construct gigabytes of test data."""
    monkeypatch.setattr(get_settings(), "docx_max_uncompressed_bytes", 1)

    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("Clause 1: Liability shall be unlimited.")
    doc.save(buf)

    with pytest.raises(DocumentParsingError, match="decompress"):
        extract_text(buf.getvalue(), DOCX_MIME)


def test_docx_within_the_decompression_limit_still_extracts_normally() -> None:
    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("Clause 1: Liability shall be unlimited.")
    doc.save(buf)

    result = extract_text(buf.getvalue(), DOCX_MIME)
    assert "Liability shall be unlimited" in "\n".join(result.pages)


def test_txt_extraction_utf8() -> None:
    result = extract_text("Contrato de prestação de serviços.".encode(), "text/plain")
    assert result.pagination_mode == "synthetic"
    assert "prestação" in result.pages[0]


def test_txt_extraction_latin1_fallback() -> None:
    data = "Contrato com caractere especial: ç".encode("latin-1")
    result = extract_text(data, "text/plain")
    assert "ç" in result.pages[0]


def test_unsupported_mime_raises() -> None:
    with pytest.raises(DocumentParsingError):
        extract_text(b"data", "application/zip")
