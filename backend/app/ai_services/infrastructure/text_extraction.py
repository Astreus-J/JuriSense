"""Text extraction for the three supported formats (`docs/05-ai/document-intelligence.md`).

Produces a list of per-page text blocks plus a `pagination_mode` flag
(`native`/`synthetic`), matching the document's "Pagination for non-paginated formats"
section: PDF pages are real, DOCX/TXT have no intrinsic page boundary so a fixed
character budget stands in for one. `text_span` provenance the AI Gateway asks the model
to produce is always relative to *one page's* text (its 0-based character offset within
that page block), never the whole-document offset -- this keeps prompts simple and
matches how a reviewer would page through the source file.

Scope decision (documented, not silent): OCR for scanned/image-only PDFs is an
explicit `TBD` in `document-intelligence.md` pending a vendor benchmark -- this module
does not implement it. A PDF page that yields no extractable native text comes back as
an empty string; `extraction_service` treats an entirely-empty document as a parsing
failure (`document.processing_failed`) rather than silently extracting nothing.
Structural/token-budget chunking for very long contracts (same document's "Chunking
strategy" section) is also out of scope for this phase -- see
`Settings.ai_extraction_max_input_chars` and its docstring in `app/core/config.py`.
"""

import io
import zipfile
from dataclasses import dataclass

from docx import Document as DocxDocument
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.ai_services.domain.errors import DocumentParsingError
from app.core.config import get_settings

# Synthetic "page" budget for DOCX/TXT, approximating a typical rendered page of a
# commercial contract (character count, not exact -- document-intelligence.md only
# requires this be applied *consistently* for a given Document, not that it match any
# particular renderer's pagination).
_SYNTHETIC_PAGE_CHARS = 3000


@dataclass(frozen=True)
class ExtractedDocument:
    pages: list[str]
    pagination_mode: str  # "native" | "synthetic"


def _synthetic_pages(full_text: str) -> list[str]:
    if not full_text:
        return [""]
    return [
        full_text[i : i + _SYNTHETIC_PAGE_CHARS]
        for i in range(0, len(full_text), _SYNTHETIC_PAGE_CHARS)
    ]


def _extract_pdf(data: bytes) -> ExtractedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, ValueError) as exc:
        raise DocumentParsingError(f"Could not parse PDF: {exc}") from exc
    if not pages:
        raise DocumentParsingError("PDF has no pages")
    return ExtractedDocument(pages=pages, pagination_mode="native")


def _reject_docx_decompression_bomb(data: bytes) -> None:
    """Security audit (2026-10-12): DOCX is a zip container -- the 25MB compressed-size
    cap (`MaxBodySizeMiddleware`) bounds the uploaded file, not what it decompresses to.
    A crafted DOCX well under that cap can still decompress to gigabytes (classic
    zip-bomb ratio), and `DocxDocument(...)` -> `python-docx` -> `zipfile` has no
    per-member/total decompressed-size cap of its own -- confirmed by reading both
    libraries' source. Mirrors `app.integrations.infrastructure.sanctions_source`'s own
    capped-download discipline for exactly this risk class, checked BEFORE
    `DocxDocument` ever reads any member's content."""
    limit = get_settings().docx_max_uncompressed_bytes
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total_uncompressed = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise DocumentParsingError(f"Could not parse DOCX: {exc}") from exc
    if total_uncompressed > limit:
        raise DocumentParsingError(
            f"DOCX would decompress to {total_uncompressed} bytes, exceeding the "
            f"{limit}-byte limit"
        )


def _extract_docx(data: bytes) -> ExtractedDocument:
    _reject_docx_decompression_bomb(data)
    try:
        doc = DocxDocument(io.BytesIO(data))
        parts: list[str] = [p.text for p in doc.paragraphs if p.text]
        # Tables flattened to text with a structural marker, per
        # document-intelligence.md, so a clause embedded in a table (e.g. an SLA
        # matrix) is not silently dropped.
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text for cell in row.cells]
                if any(cells):
                    parts.append("[TABLE ROW] " + " | ".join(cells))
    except (KeyError, ValueError, OSError) as exc:
        raise DocumentParsingError(f"Could not parse DOCX: {exc}") from exc
    full_text = "\n".join(parts)
    return ExtractedDocument(pages=_synthetic_pages(full_text), pagination_mode="synthetic")


def _extract_txt(data: bytes) -> ExtractedDocument:
    for encoding in ("utf-8", "latin-1"):
        try:
            full_text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise DocumentParsingError("Could not decode TXT file with utf-8 or latin-1")
    return ExtractedDocument(pages=_synthetic_pages(full_text), pagination_mode="synthetic")


_EXTRACTORS = {
    "application/pdf": _extract_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_docx,
    "text/plain": _extract_txt,
}


def extract_text(data: bytes, mime_type: str) -> ExtractedDocument:
    """Dispatches to the right per-format extractor. `mime_type` is the *sniffed* type
    from `app.contracts.infrastructure.document_scanner` (SEC-003), never the
    client-declared one, so this only ever sees one of the three supported values --
    any other type is rejected at upload time, before this function is ever called."""
    extractor = _EXTRACTORS.get(mime_type)
    if extractor is None:
        raise DocumentParsingError(f"Unsupported mime type for extraction: {mime_type}")
    return extractor(data)
